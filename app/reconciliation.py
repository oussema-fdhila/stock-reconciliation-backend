import io
from typing import Optional

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .plant_configs import get_plant_signes


# ============================================================
# NETTOYAGE DES DONNÉES
# ============================================================

def clean_text_series(s: pd.Series) -> pd.Series:
    return (
        s.fillna("")
        .astype(str)
        .str.replace("\x00", "", regex=False)
        .str.replace("\xa0", " ", regex=False)
        .str.replace("\ufeff", "", regex=False)
        .str.strip()
        .replace({
            "nan": "",
            "NaN": "",
            "None": ""
        })
    )


def read_bytes_without_nulls(data: bytes) -> bytes:
    return data.replace(b"\x00", b"")


# ============================================================
# DÉTECTION DU SÉPARATEUR CSV
# ============================================================

def detect_separator(data: bytes, sample_lines: int = 30) -> str:
    text = data.decode("latin1", errors="ignore")

    lines = text.splitlines()[:sample_lines]

    candidates = [",", ";", "\t", "|"]

    scores = {
        sep: sum(line.count(sep) for line in lines)
        for sep in candidates
    }

    return max(scores, key=scores.get)


# ============================================================
# LECTURE CSV
# ============================================================

def read_csv_bytes(data: bytes, skiprows: int) -> pd.DataFrame:

    data = read_bytes_without_nulls(data)

    sep = detect_separator(data)

    return pd.read_csv(
        io.BytesIO(data),
        sep=sep,
        skiprows=skiprows,
        encoding="latin1",
        dtype=str,
        low_memory=False,
    )


# ============================================================
# NETTOYAGE DATAFRAME
# ============================================================

def clean_dataframe(df: pd.DataFrame) -> pd.DataFrame:

    new_cols = []

    for col in df.columns:

        name = (
            str(col)
            .replace("\xa0", " ")
            .replace("\ufeff", "")
            .strip()
        )

        new_cols.append(name)

    df = df.copy()

    df.columns = new_cols

    for i in range(df.shape[1]):

        df.iloc[:, i] = clean_text_series(
            df.iloc[:, i]
        )

    return df


# ============================================================
# TRAITEMENT STOCK SI / SF
# ============================================================

def process_stock(data: bytes) -> pd.DataFrame:

    df = clean_dataframe(
        read_csv_bytes(
            data,
            skiprows=9
        )
    )

    if df.shape[1] <= 9:

        raise ValueError(
            "Fichier stock invalide : "
            "au moins 10 colonnes sont nécessaires."
        )

    # Colonne G = Stock
    stock = pd.to_numeric(
        clean_text_series(
            df.iloc[:, 6]
        ).str.replace(
            ",",
            ".",
            regex=False
        ),
        errors="coerce",
    )

    # Colonne J = Valeur
    value = pd.to_numeric(
        clean_text_series(
            df.iloc[:, 9]
        ).str.replace(
            ",",
            ".",
            regex=False
        ),
        errors="coerce",
    )

    # Prix unitaire
    df["PU"] = (
        value / stock
    ).replace(
        [float("inf"), -float("inf")],
        0
    ).fillna(0)

    # PartNumber
    df["_PartNumber"] = clean_text_series(
        df.iloc[:, 0]
    )

    # Stock
    df["_Stock"] = stock.fillna(0)

    return df


# ============================================================
# NORMALISATION BOOKTYPE
# ============================================================

def normalize_booktype(booktype: pd.Series) -> pd.Series:

    """
    Transforme par exemple :

    01 -> 1
    02 -> 2
    04 -> 4
    05 -> 5
    12 -> 12
    24 -> 24

    Cela permet de gérer les FMALS032 où BookType
    peut contenir des zéros en début de valeur.
    """

    return booktype.str.replace(
        r"^0+(\d+)$",
        r"\1",
        regex=True
    )


# ============================================================
# TRAITEMENT FMALS032
# ============================================================

def process_movement(
    data: bytes,
    file_number: int,
    signes: dict
):

    df = clean_dataframe(
        read_csv_bytes(
            data,
            skiprows=12
        )
    )

    if df.shape[1] <= 11:

        raise ValueError(
            f"FMALS032 #{file_number} invalide : "
            "au moins 12 colonnes sont nécessaires."
        )

    # ========================================================
    # G_Type
    # ========================================================

    col_g = df.columns[6]
    col_l = df.columns[11]

    df.insert(
        12,
        "G_Type",
        clean_text_series(df[col_g])
        + "_"
        + clean_text_series(df[col_l])
    )

    # ========================================================
    # COLONNES OBLIGATOIRES
    # ========================================================

    required = [
        "PartNr.",
        "BookType",
        "Quantity",
        "Type"
    ]

    missing = [
        c for c in required
        if c not in df.columns
    ]

    if missing:

        raise ValueError(
            f"FMALS032 #{file_number} : "
            f"colonnes manquantes : {missing}"
        )

    # ========================================================
    # SÉLECTION DES COLONNES
    # ========================================================

    m = df[
        [
            "PartNr.",
            "BookType",
            "Quantity",
            "Type",
            "G_Type"
        ]
    ].copy()

    # ========================================================
    # NETTOYAGE
    # ========================================================

    for col in [
        "PartNr.",
        "BookType",
        "Type",
        "G_Type"
    ]:

        m[col] = clean_text_series(
            m[col]
        )

    # ========================================================
    # QUANTITÉ
    # ========================================================

    m["Quantity"] = pd.to_numeric(
        clean_text_series(
            m["Quantity"]
        )
        .str.replace(
            "\xa0",
            "",
            regex=False
        )
        .str.replace(
            ",",
            ".",
            regex=False
        ),
        errors="coerce",
    ).fillna(0)

    # ========================================================
    # NORMALISATION BOOKTYPE
    # ========================================================

    m["BookType_normalise"] = normalize_booktype(
        m["BookType"]
    )

    # ========================================================
    # CODE MOUVEMENT
    #
    # Exemple :
    #
    # BookType = 01
    # Type     = RECEIPT
    #
    # devient :
    #
    # 1_RECEIPT
    # ========================================================

    m["Mouv"] = (
        m["BookType_normalise"]
        + "_"
        + m["Type"]
    )

    # ========================================================
    # APPLICATION DU SIGNE DU PLANT
    # ========================================================

    m["Signe"] = (
        m["Mouv"]
        .map(signes)
        .fillna(0)
    )

    # ========================================================
    # QUANTITÉ SIGNÉE
    # ========================================================

    m["Quantity_signed"] = (
        m["Quantity"]
        * m["Signe"]
    )

    # ========================================================
    # MOUVEMENTS INCONNUS
    # ========================================================

    unknown = sorted(
        set(
            m.loc[
                (
                    ~m["Mouv"].isin(signes)
                )
                &
                (
                    m["Mouv"] != ""
                ),
                "Mouv"
            ]
        )
    )

    # ========================================================
    # TCD
    #
    # Les colonnes correspondent aux codes mouvements.
    # Les valeurs sont les quantités.
    # ========================================================

    tcd = pd.pivot_table(
        m,
        index="PartNr.",
        columns="Mouv",
        values="Quantity",
        aggfunc="sum",
        fill_value=0,
    ).reset_index()

    tcd = tcd.rename(
        columns={
            "PartNr.": "PartNumber"
        }
    )

    # ========================================================
    # APPLICATION DES SIGNES AU TCD
    # ========================================================

    movement_cols = [
        c
        for c in tcd.columns
        if c != "PartNumber"
    ]

    for col in movement_cols:

        sign = signes.get(
            str(col).strip(),
            0
        )

        tcd[col] = (
            pd.to_numeric(
                tcd[col],
                errors="coerce"
            )
            .fillna(0)
            * sign
        )

    # ========================================================
    # TOTAL DU TCD
    # ========================================================

    if movement_cols:

        tcd["Total"] = tcd[
            movement_cols
        ].sum(axis=1)

    else:

        tcd["Total"] = 0

    return (
        tcd,
        len(unknown),
        unknown
    )


# ============================================================
# CONSTRUCTION DE LA RÉCONCILIATION
# ============================================================

def build_reconciliation(
    si: pd.DataFrame,
    sf: pd.DataFrame,
    tcd1: pd.DataFrame,
    tcd2: pd.DataFrame
) -> pd.DataFrame:

    # ========================================================
    # PART NUMBERS
    # ========================================================

    parts_si = (
        si.iloc[:, 0]
        .astype(str)
        .str.strip()
    )

    parts_sf = (
        sf.iloc[:, 0]
        .astype(str)
        .str.strip()
    )

    parts_si = parts_si[
        parts_si != ""
    ]

    parts_sf = parts_sf[
        parts_sf != ""
    ]

    parts = sorted(
        set(parts_si)
        |
        set(parts_sf)
    )

    recon = pd.DataFrame({
        "PartNumber": parts
    })

    # ========================================================
    # STOCK SI
    # ========================================================

    si_par = (
        si.groupby(
            "_PartNumber"
        )["_Stock"]
        .sum()
    )

    # ========================================================
    # STOCK SF
    # ========================================================

    sf_par = (
        sf.groupby(
            "_PartNumber"
        )["_Stock"]
        .sum()
    )

    # ========================================================
    # MOUVEMENT 1
    # ========================================================

    m1 = (
        tcd1
        .set_index("PartNumber")["Total"]
    )

    # ========================================================
    # MOUVEMENT 2
    # ========================================================

    m2 = (
        tcd2
        .set_index("PartNumber")["Total"]
    )

    # ========================================================
    # MAPPING
    # ========================================================

    recon["SI"] = (
        recon["PartNumber"]
        .map(si_par)
        .fillna(0)
    )

    recon["SF"] = (
        recon["PartNumber"]
        .map(sf_par)
        .fillna(0)
    )

    recon["Mouvement 1"] = (
        recon["PartNumber"]
        .map(m1)
        .fillna(0)
    )

    recon["Mouvement 2"] = (
        recon["PartNumber"]
        .map(m2)
        .fillna(0)
    )

    # ========================================================
    # MOUVEMENT TOTAL
    # ========================================================

    recon["Mouvement"] = (
        recon["Mouvement 1"]
        +
        recon["Mouvement 2"]
    )

    # ========================================================
    # TYPE
    #
    # SF prioritaire, sinon SI
    # ========================================================

    type_si = (
        si
        .drop_duplicates(
            "_PartNumber"
        )
        .set_index(
            "_PartNumber"
        )
        .iloc[:, 3]
    )

    type_sf = (
        sf
        .drop_duplicates(
            "_PartNumber"
        )
        .set_index(
            "_PartNumber"
        )
        .iloc[:, 3]
    )

    recon["Type"] = (
        recon["PartNumber"]
        .map(type_sf)
        .fillna(
            recon["PartNumber"]
            .map(type_si)
        )
        .fillna("")
    )

    # ========================================================
    # PRIX UNITAIRE
    #
    # SF prioritaire, sinon SI
    # ========================================================

    pu_si = (
        si
        .drop_duplicates(
            "_PartNumber"
        )
        .set_index(
            "_PartNumber"
        )["PU"]
    )

    pu_sf = (
        sf
        .drop_duplicates(
            "_PartNumber"
        )
        .set_index(
            "_PartNumber"
        )["PU"]
    )

    recon["PU"] = (
        recon["PartNumber"]
        .map(pu_sf)
        .fillna(
            recon["PartNumber"]
            .map(pu_si)
        )
        .fillna(0)
    )

    recon["PU"] = pd.to_numeric(
        recon["PU"],
        errors="coerce"
    ).fillna(0)

    # ========================================================
    # ÉCART
    #
    # Ecart = SF - SI - Mouvement
    # ========================================================

    recon["Ecart"] = (
        recon["SF"]
        -
        recon["SI"]
        -
        recon["Mouvement"]
    )

    # ========================================================
    # ÉCART EN EURO
    # ========================================================

    recon["Ecart €"] = (
        recon["Ecart"]
        *
        recon["PU"]
    )

    # ========================================================
    # ORDRE FINAL DES COLONNES
    # ========================================================

    return recon[
        [
            "PartNumber",
            "Type",
            "PU",
            "SI",
            "Mouvement 1",
            "Mouvement 2",
            "Mouvement",
            "SF",
            "Ecart",
            "Ecart €"
        ]
    ]


# ============================================================
# EXPORT EXCEL
# ============================================================

def export_excel(
    recon: pd.DataFrame,
    tcd1: pd.DataFrame,
    tcd2: pd.DataFrame
) -> bytes:

    out = io.BytesIO()

    with pd.ExcelWriter(
        out,
        engine="openpyxl"
    ) as writer:

        recon.to_excel(
            writer,
            sheet_name="Réconcil",
            index=False
        )

        tcd1.to_excel(
            writer,
            sheet_name="TCD1",
            index=False
        )

        tcd2.to_excel(
            writer,
            sheet_name="TCD2",
            index=False
        )

    out.seek(0)

    # ========================================================
    # FORMATAGE EXCEL
    # ========================================================

    wb = load_workbook(out)

    for ws in wb.worksheets:

        ws.freeze_panes = "A2"

        ws.auto_filter.ref = ws.dimensions

        # Header
        for cell in ws[1]:

            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )

            cell.fill = PatternFill(
                "solid",
                fgColor="002857"
            )

            cell.alignment = Alignment(
                horizontal="center"
            )

        # Largeur colonnes
        for col_cells in ws.columns:

            letter = get_column_letter(
                col_cells[0].column
            )

            max_len = min(
                max(
                    len(
                        str(
                            c.value
                            or ""
                        )
                    )
                    for c in col_cells
                ) + 2,
                45
            )

            ws.column_dimensions[
                letter
            ].width = max(
                10,
                max_len
            )

        # Format numérique Réconcil
        if ws.title == "Réconcil":

            for col in [
                3,
                4,
                5,
                6,
                7,
                8,
                9,
                10
            ]:

                for cell in ws[
                    get_column_letter(col)
                ][1:]:

                    cell.number_format = (
                        '#,##0.00'
                    )

    final = io.BytesIO()

    wb.save(final)

    final.seek(0)

    return final.getvalue()


# ============================================================
# FONCTION PRINCIPALE DE RÉCONCILIATION
# ============================================================

def reconcile(
    si_bytes: bytes,
    sf_bytes: bytes,
    movement1_bytes: bytes,
    movement2_bytes: Optional[bytes] = None,
    plant: str = "48"
):

    # ========================================================
    # RÉCUPÉRATION DES SIGNES DU PLANT
    # ========================================================

    plant = str(
        plant
    ).strip()

    signes = get_plant_signes(
        plant
    )

    # ========================================================
    # TRAITEMENT SI / SF
    # ========================================================

    si = process_stock(
        si_bytes
    )

    sf = process_stock(
        sf_bytes
    )

    # ========================================================
    # FMALS032 1
    # ========================================================

    tcd1, unknown1_count, unknown1 = (
        process_movement(
            movement1_bytes,
            1,
            signes
        )
    )

    # ========================================================
    # FMALS032 2
    # ========================================================

    if movement2_bytes:

        tcd2, unknown2_count, unknown2 = (
            process_movement(
                movement2_bytes,
                2,
                signes
            )
        )

    else:

        tcd2 = pd.DataFrame(
            columns=[
                "PartNumber",
                "Total"
            ]
        )

        unknown2_count = 0

        unknown2 = []

    # ========================================================
    # RÉCONCILIATION
    # ========================================================

    recon = build_reconciliation(
        si,
        sf,
        tcd1,
        tcd2
    )

    # ========================================================
    # EXPORT EXCEL
    # ========================================================

    excel = export_excel(
        recon,
        tcd1,
        tcd2
    )

    # ========================================================
    # SUMMARY
    # ========================================================

    summary = {

        "plant": plant,

        "partnumbers": len(
            recon
        ),

        "tcd1_partnumbers": len(
            tcd1
        ),

        "tcd2_partnumbers": len(
            tcd2
        ),

        "si_total": float(
            recon["SI"].sum()
        ),

        "sf_total": float(
            recon["SF"].sum()
        ),

        "movement1_total": float(
            recon["Mouvement 1"].sum()
        ),

        "movement2_total": float(
            recon["Mouvement 2"].sum()
        ),

        "movement_total": float(
            recon["Mouvement"].sum()
        ),

        "ecart_total": float(
            recon["Ecart"].sum()
        ),

        "ecart_eur_total": float(
            recon["Ecart €"].sum()
        ),

        "unknown_movements_file_1":
            unknown1_count,

        "unknown_movements_file_2":
            unknown2_count,

        "unknown_movement_codes_file_1":
            unknown1,

        "unknown_movement_codes_file_2":
            unknown2,
    }

    # ========================================================
    # PREVIEW
    # ========================================================

    preview_df = recon.head(
        100
    ).where(
        pd.notnull(
            recon.head(100)
        ),
        None
    )

    preview = preview_df.to_dict(
        orient="records"
    )

    return (
        excel,
        summary,
        preview
    )
